from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
presentation = Presentation()
presentation.slide_width = Inches(13.333)
presentation.slide_height = Inches(7.5)


def slide(title: str, bullets: list[str], image: str | None = None) -> None:
    layout = presentation.slide_layouts[5]
    page = presentation.slides.add_slide(layout)
    page.shapes.title.text = title
    title_shape = page.shapes.title
    title_shape.text_frame.paragraphs[0].font.size = Pt(28)
    box = page.shapes.add_textbox(Inches(.8), Inches(1.4), Inches(6.1 if image else 11.7), Inches(5.3))
    frame = box.text_frame
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.font.size = Pt(20)
        paragraph.space_after = Pt(12)
    if image:
        page.shapes.add_picture(str(ROOT / image), Inches(7.2), Inches(1.5), width=Inches(5.5))


slide("Projekt 4 — Artificial Bee Colony", ["Optymalizacja funkcji wielu zmiennych z MealPy", "Porównanie z algorytmami genetycznymi P1/P2"])
slide("Inspiracja biologiczna", ["Kolonia pszczół poszukuje źródeł pożywienia.", "Źródło pożywienia reprezentuje jedno rozwiązanie.", "Jakość źródła odpowiada wartości funkcji celu."])
slide("Trzy typy pszczół", ["Pracujące: lokalnie poprawiają własne źródła.", "Obserwujące: wybierają dobre źródła na podstawie jakości.", "Zwiadowcze: losują nowe źródła po przekroczeniu limitu prób."])
slide("Pseudokod ABC", ["1. Utwórz populację źródeł.", "2. Wykonuj lokalne przeszukiwanie i selekcję lepszych źródeł.", "3. Porzucone źródła zastępuj losowymi zwiadowcami.", "4. Zapamiętaj najlepsze rozwiązanie."])
slide("Przykład analityczny", ["Źródło x=(2,4), sąsiad x'=(2.5,3.5).", "Obliczamy f(x) i f(x').", "Jeśli f(x') < f(x), sąsiad zastępuje źródło.", "Brak poprawy przez n_limits uruchamia pszczołę zwiadowczą."])
slide("Funkcje testowe", ["Ackley 01", "Hyperellipsoid", "CEC2014 F1 i F3", "Funkcja kwadratowa"])
slide("Konfiguracja eksperymentu", ["Populacja: 20 osobników", "50 epok, n_limits=15", "5 niezależnych powtórzeń", "Stałe seedy dla porównywalności"])
slide("Zbieżność Ackley", ["ABC i GA porównano przy tej samej liczbie epok i osobników."], "convergence_ackley01.png")
slide("Zbieżność CEC2014 F3", ["Funkcja wielowymiarowa o trudnym krajobrazie celu."], "convergence_cec2014_f3.png")
slide("ABC kontra P1/P2", ["Porównanie wykonano na wspólnych funkcjach i seedach.", "Uwzględniono jakość rozwiązania oraz czas obliczeń."], "comparison_p1_p2.png")
slide("Wpływ konfiguracji", ["Zmieniano populację, liczbę epok i n_limits.", "Większa liczba epok zwykle poprawiała wynik kosztem czasu."], "configurations_abc.png")
slide("Wnioski", ["ABC łączy eksplorację dzięki zwiadowcom z lokalnym przeszukiwaniem.", "Na części benchmarków uzyskał lepsze wyniki niż GA.", "GA bywa szybszy przy tej samej liczbie epok.", "Najlepsza konfiguracja zależy od funkcji celu."])
slide("Dziękuję za uwagę", ["Kod, wyniki, sprawozdanie i prezentacja znajdują się w katalogu sprawozdanie_p4."])

presentation.save(ROOT / "P4_PREZENTACJA.pptx")
print(ROOT / "P4_PREZENTACJA.pptx")
